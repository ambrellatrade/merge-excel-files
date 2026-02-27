from __future__ import annotations

import argparse
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

try:
    import xlrd  # type: ignore
except ImportError:  # optional dependency for .xls support
    xlrd = None

KNOWN_TYPES = {
    "Детали": "DETAL",
    "Сборочные единицы": "SBED",
    "Стандартные изделия": "STAND_IZD",
    "Прочие изделия": "PROCH_IZD",
    "Комплекты": "KOMPLEKT",
    "Материалы": "MATERIAL",
}

BRANCH_TYPES = {"SBED", "KOMPLEKT"}
OUTPUT_COLUMNS = ["Type", "Code", "Name", "Kol", "Vhod", "Pos"]


@dataclass
class SpecRow:
    file_code: str
    file_path: Path
    pos: Any
    code: str
    name: str
    quantities: list[float]
    type_text: str
    type_code: str


@dataclass
class HeadSpec:
    base_code: str
    name: str
    path: Path


@dataclass
class BuildContext:
    missing_spec_files: list[dict[str, Any]]
    unknown_types: list[dict[str, Any]]
    rows_without_qty: list[dict[str, Any]]
    parse_errors: list[dict[str, Any]]
    duplicates_by_code: list[dict[str, Any]]
    cycles: list[dict[str, Any]]
    output_rows: list[dict[str, Any]]
    files_parsed: int
    heads_count: int
    executions_count: int


def parse_file_name(path: Path) -> tuple[str, str]:
    stem = path.stem.strip()
    parts = stem.split(maxsplit=1)
    if not parts:
        raise ValueError(f"Cannot parse file name: {path.name}")
    code = parts[0].strip()
    name = parts[1].strip() if len(parts) > 1 else code
    return code, name


def execution_code(base_code: str, k: int) -> str:
    match = re.match(r"^(.*?)-(\d{2})$", base_code)
    if match:
        prefix, num = match.groups()
        return f"{prefix}-{int(num) + k:02d}"
    if k == 0:
        return base_code
    return f"{base_code}-{k:02d}"


def normalize_qty(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def try_int(value: Any) -> int | None:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def read_sheet_rows(path: Path) -> list[list[Any]]:
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Для чтения .xlsx установите зависимость: pip install openpyxl") from exc
        wb = load_workbook(path, data_only=True, read_only=True)
        best_ws = max(wb.worksheets, key=lambda ws: ws.max_row)
        return [list(row) for row in best_ws.iter_rows(values_only=True)]
    if path.suffix.lower() == ".xls":
        if xlrd is None:
            raise RuntimeError("Для чтения .xls установите зависимость: pip install xlrd==2.0.1")
        book = xlrd.open_workbook(path)
        sheet = max((book.sheet_by_index(i) for i in range(book.nsheets)), key=lambda s: s.nrows)
        return [sheet.row_values(i) for i in range(sheet.nrows)]
    raise RuntimeError(f"Unsupported file extension: {path.suffix}")


def detect_data_start(rows: list[list[Any]]) -> int:
    for idx, row in enumerate(rows):
        if len(row) < 4:
            continue
        pos = try_int(row[0])
        name = str(row[2]).strip() if row[2] is not None else ""
        type_text = str(row[-1]).strip() if row[-1] is not None else ""
        if pos is not None and pos >= 1 and name and type_text:
            return idx
    return 0


def parse_spec(path: Path, ctx: BuildContext) -> list[SpecRow]:
    code, _ = parse_file_name(path)
    rows = read_sheet_rows(path)
    start = detect_data_start(rows)
    parsed: list[SpecRow] = []
    for raw in rows[start:]:
        if not raw or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        if len(raw) < 4:
            continue
        pos = try_int(raw[0])
        if pos is None:
            continue
        item_code = "" if raw[1] is None else str(raw[1]).strip()
        name = "" if raw[2] is None else str(raw[2]).strip()
        type_text = "" if raw[-1] is None else str(raw[-1]).strip()
        type_code = KNOWN_TYPES.get(type_text, "DOCUMENT")
        if type_text not in KNOWN_TYPES:
            ctx.unknown_types.append(
                {
                    "file": str(path),
                    "Pos": pos,
                    "Code": item_code,
                    "Name": name,
                    "TypeText": type_text,
                }
            )
        qty_raw = raw[3:-1]
        quantities = [normalize_qty(v) for v in qty_raw] if qty_raw else [0.0]
        parsed.append(
            SpecRow(
                file_code=code,
                file_path=path,
                pos=pos,
                code=item_code,
                name=name,
                quantities=quantities,
                type_text=type_text,
                type_code=type_code,
            )
        )
    return sorted(parsed, key=lambda r: (r.pos if isinstance(r.pos, int) else 10**9))


def build_spec_index(specs_dir: Path, ctx: BuildContext) -> dict[str, Path]:
    index: dict[str, Path] = {}
    duplicates: dict[str, list[Path]] = {}
    for path in sorted(specs_dir.rglob("*")):
        if path.suffix.lower() not in {".xls", ".xlsx"}:
            continue
        try:
            code, _ = parse_file_name(path)
        except Exception:
            continue
        if code in index:
            duplicates.setdefault(code, [index[code]]).append(path)
            continue
        index[code] = path
    for code, paths in duplicates.items():
        ctx.duplicates_by_code.append({"Code": code, "Files": "\n".join(str(p) for p in paths)})
    return index


def load_heads(base: Path, head_dir: Path) -> list[HeadSpec]:
    heads: list[HeadSpec] = []
    for file in sorted(head_dir.glob("*")):
        if file.suffix.lower() not in {".xls", ".xlsx"}:
            continue
        code, name = parse_file_name(file)
        heads.append(HeadSpec(base_code=code, name=name, path=file))
    if heads:
        return heads

    roots_path = base / "input" / "roots.txt"
    if roots_path.exists():
        for line in roots_path.read_text(encoding="utf-8").splitlines():
            code = line.strip()
            if code:
                heads.append(HeadSpec(base_code=code, name=code, path=Path("")))
    if not heads:
        raise RuntimeError("Не найдены головные спецификации в input/head и отсутствует input/roots.txt")
    return heads


def choose_head_path(head: HeadSpec, spec_index: dict[str, Path]) -> Path:
    if head.path and head.path.exists():
        return head.path
    if head.base_code in spec_index:
        return spec_index[head.base_code]
    raise RuntimeError(f"Не найден файл спецификации для головы: {head.base_code}")


def add_output_row(ctx: BuildContext, type_code: str, code: str, name: str, kol: float | int, vhod: str, pos: Any) -> None:
    ctx.output_rows.append(
        {
            "Type": type_code,
            "Code": code or "",
            "Name": name or "",
            "Kol": kol if kol is not None else "",
            "Vhod": vhod or "",
            "Pos": pos if pos is not None else "",
        }
    )


def expand_spec(
    *,
    spec_path: Path,
    parent_exec_code: str,
    k: int,
    spec_index: dict[str, Path],
    ctx: BuildContext,
    stack: list[tuple[str, str]],
) -> None:
    try:
        rows = parse_spec(spec_path, ctx)
        ctx.files_parsed += 1
    except Exception as exc:
        ctx.parse_errors.append({"path": str(spec_path), "error": str(exc)})
        return

    for row in rows:
        qty = row.quantities[k] if k < len(row.quantities) else 0.0
        if qty <= 0:
            ctx.rows_without_qty.append(
                {
                    "file": str(row.file_path),
                    "Pos": row.pos,
                    "Code": row.code,
                    "Name": row.name,
                    "execution": execution_code(parse_file_name(spec_path)[0], k),
                }
            )
            continue
        add_output_row(ctx, row.type_code, row.code, row.name, qty, parent_exec_code, row.pos)

        if row.type_code not in BRANCH_TYPES:
            continue
        if not row.code.startswith("КУМП."):
            continue
        if row.code not in spec_index:
            ctx.missing_spec_files.append(
                {
                    "parent": str(spec_path),
                    "child_code": row.code,
                    "Pos": row.pos,
                    "Name": row.name,
                    "Type": row.type_code,
                }
            )
            continue

        child_exec_code = execution_code(row.code, k)
        node = (row.code, child_exec_code)
        if node in stack:
            ctx.cycles.append(
                {
                    "cycle_at": row.code,
                    "execution": child_exec_code,
                    "stack": " -> ".join([f"{c}:{e}" for c, e in stack] + [f"{row.code}:{child_exec_code}"]),
                }
            )
            continue

        stack.append(node)
        expand_spec(
            spec_path=spec_index[row.code],
            parent_exec_code=child_exec_code,
            k=k,
            spec_index=spec_index,
            ctx=ctx,
            stack=stack,
        )
        stack.pop()


def write_output_from_template(template_path: Path, output_path: Path, rows: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("Для записи Excel установите зависимость: pip install openpyxl") from exc

    if template_path.exists():
        wb = load_workbook(template_path)
        ws = wb.active
        ws.delete_rows(2, ws.max_row)
    else:
        wb = Workbook()
        ws = wb.active
        for idx, col in enumerate(OUTPUT_COLUMNS, start=1):
            ws.cell(row=1, column=idx, value=col)
    header = [ws.cell(row=1, column=i).value for i in range(1, len(OUTPUT_COLUMNS) + 1)]
    if header != OUTPUT_COLUMNS:
        for idx, col in enumerate(OUTPUT_COLUMNS, start=1):
            ws.cell(row=1, column=idx, value=col)

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, key in enumerate(OUTPUT_COLUMNS, start=1):
            value = row.get(key, "")
            if value is None:
                value = ""
            ws.cell(row=r_idx, column=c_idx, value=value)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_report(report_path: Path, ctx: BuildContext, spec_count: int) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("Для записи Excel установите зависимость: pip install openpyxl") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    summary_rows = [
        ("files_found_in_specs", spec_count),
        ("files_parsed", ctx.files_parsed),
        ("rows_exported", len(ctx.output_rows)),
        ("heads_count", ctx.heads_count),
        ("executions_count", ctx.executions_count),
    ]
    for i, (k, v) in enumerate(summary_rows, start=1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    add_sheet(wb, "missing_spec_files", ["parent", "child_code", "Pos", "Name", "Type"], ctx.missing_spec_files)
    add_sheet(wb, "unknown_types", ["file", "Pos", "Code", "Name", "TypeText"], ctx.unknown_types)
    add_sheet(wb, "rows_without_qty", ["file", "Pos", "Code", "Name", "execution"], ctx.rows_without_qty)
    add_sheet(wb, "parse_errors", ["path", "error"], ctx.parse_errors)
    add_sheet(wb, "duplicates_by_code", ["Code", "Files"], ctx.duplicates_by_code)
    add_sheet(wb, "cycles", ["cycle_at", "execution", "stack"], ctx.cycles)

    report_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(report_path)


def add_sheet(wb: Any, name: str, columns: list[str], rows: Iterable[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=name)
    for col_idx, key in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=key)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(columns, start=1):
            value = row.get(key, "")
            if value is None:
                value = ""
            ws.cell(row=row_idx, column=col_idx, value=value)


def ensure_structure(base: Path) -> None:
    (base / "input" / "head").mkdir(parents=True, exist_ok=True)
    (base / "input" / "specs").mkdir(parents=True, exist_ok=True)
    (base / "template").mkdir(parents=True, exist_ok=True)
    (base / "output").mkdir(parents=True, exist_ok=True)
    (base / "logs").mkdir(parents=True, exist_ok=True)


def run(args: argparse.Namespace) -> int:
    base = Path(args.base)
    ensure_structure(base)

    log_path = base / "logs" / "run.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8"), logging.StreamHandler()],
    )

    ctx = BuildContext(
        missing_spec_files=[],
        unknown_types=[],
        rows_without_qty=[],
        parse_errors=[],
        duplicates_by_code=[],
        cycles=[],
        output_rows=[],
        files_parsed=0,
        heads_count=0,
        executions_count=0,
    )

    specs_dir = base / normalize_rel_path(args.specs_dir)
    head_dir = base / normalize_rel_path(args.head_dir)
    template = base / normalize_rel_path(args.template)
    out_path = base / normalize_rel_path(args.out)
    report_path = base / normalize_rel_path(args.report)

    spec_index = build_spec_index(specs_dir, ctx)
    heads = load_heads(base, head_dir)
    ctx.heads_count = len(heads)

    for head in heads:
        try:
            head_path = choose_head_path(head, spec_index)
            head_rows = parse_spec(head_path, ctx)
            ctx.files_parsed += 1
        except Exception as exc:
            ctx.parse_errors.append({"path": str(head.path or head.base_code), "error": str(exc)})
            continue

        max_exec = max((len(r.quantities) for r in head_rows), default=1)
        for k in range(max_exec):
            exec_code = execution_code(head.base_code, k)
            ctx.executions_count += 1
            add_output_row(ctx, "SBED", exec_code, head.name, 1, "", "")
            stack = [(head.base_code, exec_code)]
            expand_spec(
                spec_path=head_path,
                parent_exec_code=exec_code,
                k=k,
                spec_index=spec_index,
                ctx=ctx,
                stack=stack,
            )

    write_output_from_template(template, out_path, ctx.output_rows)
    write_report(report_path, ctx, len(spec_index))
    logging.info("Done. Rows exported: %s", len(ctx.output_rows))
    return 0


def normalize_rel_path(raw: str) -> Path:
    return Path(raw.replace("\\", "/"))
